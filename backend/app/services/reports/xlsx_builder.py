import io
import os
from typing import Optional
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .base_report import ReportData

from openpyxl.drawing.image import Image as OpenPyXLImage

def build_xlsx_report(data: ReportData) -> bytes:
    """Builds a formatted, styled executive Excel spreadsheet report (.xlsx)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório Oficial"
    ws.views.sheetView[0].showGridLines = True

    # Palettes
    color_navy = "1E293B"
    color_slate_header = "334155"
    color_light_gray = "F8FAFC"
    color_border = "CBD5E1"
    color_meta_bg = "F8FAFC"

    thin_border = Border(
        left=Side(style="thin", color=color_border),
        right=Side(style="thin", color=color_border),
        top=Side(style="thin", color=color_border),
        bottom=Side(style="thin", color=color_border)
    )

    font_title_entity = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
    font_sub_entity = Font(name="Segoe UI", size=9.5, color="334155")
    font_school = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    font_doc_title = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_meta_label = Font(name="Segoe UI", size=8.5, bold=True, color="0F172A")
    font_meta_val = Font(name="Segoe UI", size=8.5, color="1E293B")
    font_th = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
    font_td = Font(name="Segoe UI", size=9, color="0F172A")
    font_td_bold = Font(name="Segoe UI", size=9, bold=True, color="0F172A")
    font_footer = Font(name="Segoe UI", size=8, italic=True, color="64748B")

    max_sec_cols = max((len(sec.columns) for sec in data.sections), default=0)
    num_cols = max(5, len(data.columns), max_sec_cols)

    # 1. INSTITUTIONAL HEADER WITH LOGO
    has_logo = False
    if data.metadata.logo_path and os.path.exists(data.metadata.logo_path):
        try:
            img = OpenPyXLImage(data.metadata.logo_path)
            img.width = 52
            img.height = 52
            ws.merge_cells(start_row=1, start_column=1, end_row=3, end_column=1)
            ws.add_image(img, "A1")
            has_logo = True
            ws.column_dimensions["A"].width = 12
        except Exception as e:
            has_logo = False

    start_col = 2 if has_logo else 1
    ws.row_dimensions[1].height = 19
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 18

    # Merge entity lines across remaining columns for seamless presentation
    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=num_cols)
    c_pref = ws.cell(row=1, column=start_col, value=data.metadata.prefeitura.upper())
    c_pref.font = font_title_entity
    c_pref.alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=num_cols)
    c_sec = ws.cell(row=2, column=start_col, value=data.metadata.secretaria.upper())
    c_sec.font = font_sub_entity
    c_sec.alignment = Alignment(horizontal="left", vertical="center")

    school_name_display = data.metadata.school_name.upper() if data.metadata.school_name else "REDE MUNICIPAL DE ENSINO"
    if data.metadata.inep_code:
        school_name_display += f" — INEP: {data.metadata.inep_code}"
    ws.merge_cells(start_row=3, start_column=start_col, end_row=3, end_column=num_cols)
    c_esc = ws.cell(row=3, column=start_col, value=school_name_display)
    c_esc.font = font_school
    c_esc.alignment = Alignment(horizontal="left", vertical="center")

    # Spacer
    ws.row_dimensions[4].height = 6

    # 2. DOCUMENT TITLE BANNER
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=num_cols)
    cell_t = ws.cell(row=5, column=1, value=data.title.upper())
    cell_t.font = font_doc_title
    cell_t.fill = PatternFill(start_color=color_navy, end_color=color_navy, fill_type="solid")
    cell_t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[5].height = 24

    # Spacer
    ws.row_dimensions[6].height = 4

    # 3. METADATA GRID (Dynamic: only fields with real data)
    m = data.metadata
    meta_items = []
    
    if m.classroom_name and m.classroom_name.strip() not in ("", "-", "None"):
        meta_items.append(("Turma", m.classroom_name.strip()))
    
    if m.grade_year and m.grade_year.strip() not in ("", "-", "None", "Todos os Anos"):
        meta_items.append(("Série/Ano", m.grade_year.strip()))
        
    if m.shift and m.shift.strip() not in ("", "-", "None") and "( )" not in m.shift:
        meta_items.append(("Turno", m.shift.strip()))
    elif m.shift and m.shift.strip() not in ("", "-", "None") and m.classroom_name and m.classroom_name.strip() not in ("", "-", "None"):
        meta_items.append(("Turno", m.shift.strip()))
        
    if m.exam_title and m.exam_title.strip() not in ("", "-", "None"):
        meta_items.append(("Avaliação", m.exam_title.strip()))
        
    ano_val = m.school_year or str(datetime.now().year)
    meta_items.append(("Ano Letivo", ano_val))

    current_row = 7
    if meta_items:
        items_per_row = 2 if len(meta_items) == 2 else min(3, len(meta_items))
        chunks = [meta_items[i:i + items_per_row] for i in range(0, len(meta_items), items_per_row)]
        
        for chunk in chunks:
            ws.row_dimensions[current_row].height = 20
            n_items = len(chunk)
            cols_per_item = num_cols // n_items
            rem = num_cols % n_items
            
            c_start = 1
            for item_idx, (lbl, val) in enumerate(chunk):
                span = cols_per_item + (1 if item_idx < rem else 0)
                c_end = c_start + span - 1
                
                for col_i in range(c_start, c_end + 1):
                    cell_bg = ws.cell(row=current_row, column=col_i)
                    cell_bg.fill = PatternFill(start_color=color_meta_bg, end_color=color_meta_bg, fill_type="solid")
                    cell_bg.border = thin_border
                    if not cell_bg.value:
                        cell_bg.value = ""
                
                if c_end > c_start:
                    ws.merge_cells(start_row=current_row, start_column=c_start, end_row=current_row, end_column=c_end)
                    
                c_val = ws.cell(row=current_row, column=c_start, value=f"{lbl}: {val}")
                c_val.font = font_meta_label
                c_val.alignment = Alignment(horizontal="left", vertical="center")
                
                c_start = c_end + 1
            current_row += 1

    # Spacer
    ws.row_dimensions[current_row].height = 6
    current_row += 1

    # 4. SUMMARY CARDS (KPIs - beautifully distributed across columns without cutoff)
    if data.summary_cards:
        cards_per_row = min(4, num_cols)
        card_chunks = [data.summary_cards[i:i + cards_per_row] for i in range(0, len(data.summary_cards), cards_per_row)]
        for chunk in card_chunks:
            n_cards = len(chunk)
            ws.row_dimensions[current_row].height = 24

            if num_cols == 6 and n_cards == 4:
                spans = [(1, 1), (2, 3), (4, 5), (6, 6)]
            elif num_cols == 5 and n_cards == 4:
                spans = [(1, 2), (3, 3), (4, 4), (5, 5)]
            elif num_cols == 7 and n_cards == 4:
                spans = [(1, 2), (3, 4), (5, 6), (7, 7)]
            elif num_cols == 8 and n_cards == 4:
                spans = [(1, 2), (3, 4), (5, 6), (7, 8)]
            else:
                spans = []
                base = max(1, num_cols // n_cards)
                rem = num_cols % n_cards
                curr = 1
                for i in range(n_cards):
                    sz = base + (1 if i < rem else 0)
                    end_c = min(num_cols, curr + sz - 1)
                    spans.append((curr, end_c))
                    curr = end_c + 1

            for idx, (c_start, c_end) in enumerate(spans):
                if idx >= len(chunk):
                    break
                card = chunk[idx]

                for c_i in range(c_start, c_end + 1):
                    cell_k = ws.cell(row=current_row, column=c_i)
                    cell_k.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
                    cell_k.border = thin_border
                    if not cell_k.value:
                        cell_k.value = ""

                if c_end > c_start:
                    ws.merge_cells(start_row=current_row, start_column=c_start, end_row=current_row, end_column=c_end)

                cell_card = ws.cell(row=current_row, column=c_start, value=f"{card.get('label')}: {card.get('value')}")
                cell_card.font = font_meta_label
                cell_card.alignment = Alignment(horizontal="center", vertical="center")

            current_row += 1

        # Spacer before table
        ws.row_dimensions[current_row].height = 6
        current_row += 1

    # 5. DATA TABLES (Supports single table or multiple sections)
    def render_excel_table(cols, rows, start_r, sec_title=None, sec_subtitle=None):
        curr = start_r
        if sec_title:
            ws.merge_cells(start_row=curr, start_column=1, end_row=curr, end_column=num_cols)
            cell_sec = ws.cell(row=curr, column=1, value=sec_title.upper())
            cell_sec.font = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
            cell_sec.fill = PatternFill(start_color=color_navy, end_color=color_navy, fill_type="solid")
            cell_sec.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[curr].height = 22
            curr += 1
            
        t_header_row = curr
        ws.row_dimensions[t_header_row].height = 20
        for c_idx, col in enumerate(cols, start=1):
            c = ws.cell(row=t_header_row, column=c_idx, value=col.header)
            c.font = font_th
            c.fill = PatternFill(start_color=color_slate_header, end_color=color_slate_header, fill_type="solid")
            c.alignment = Alignment(horizontal=col.align, vertical="center")
            c.border = thin_border
        curr += 1

        d_start_row = curr
        for r_idx, row_dict in enumerate(rows, start=d_start_row):
            ws.row_dimensions[r_idx].height = 19
            is_even = (r_idx % 2 == 0)
            row_fill = PatternFill(start_color=color_light_gray, end_color=color_light_gray, fill_type="solid") if is_even else PatternFill(fill_type=None)

            for c_idx, col in enumerate(cols, start=1):
                val = row_dict.get(col.key)
                cell_d = ws.cell(row=r_idx, column=c_idx, value=val)
                cell_d.font = font_td
                cell_d.alignment = Alignment(horizontal=col.align, vertical="center")
                cell_d.border = thin_border
                if is_even:
                    cell_d.fill = row_fill

                # Numeric formatting
                if col.format_type == "decimal" and isinstance(val, (int, float)):
                    cell_d.number_format = "0.0"
                elif col.format_type == "percent" and isinstance(val, (int, float)):
                    cell_d.number_format = "0.0%"
                elif col.is_numeric and isinstance(val, (int, float)):
                    cell_d.number_format = "#,##0.0"
            curr += 1

        # Adjust column widths
        for col_idx in range(1, len(cols) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = len(str(cols[col_idx - 1].header))
            for r in range(t_header_row, curr):
                val = ws.cell(row=r, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))
            current_w = ws.column_dimensions[col_letter].width or 0
            ws.column_dimensions[col_letter].width = max(current_w, 14, min(max_len + 4, 48))

        return curr

    if data.sections:
        for sec in data.sections:
            current_row = render_excel_table(sec.columns, sec.rows, current_row, sec_title=sec.title, sec_subtitle=sec.subtitle)
            ws.row_dimensions[current_row].height = 6
            current_row += 1
    else:
        current_row = render_excel_table(data.columns, data.rows, current_row)

    # FOOTER SIGNATURES (Centered and cleanly merged)
    sig_row = current_row + 2
    ws.row_dimensions[sig_row].height = 18
    ws.row_dimensions[sig_row + 1].height = 16

    if len(data.signatures) >= 2:
        mid_col = num_cols // 2
        # Signature 1 (Left Half)
        for c in range(1, mid_col + 1):
            ws.cell(row=sig_row, column=c)
            ws.cell(row=sig_row + 1, column=c)
        ws.merge_cells(start_row=sig_row, start_column=1, end_row=sig_row, end_column=mid_col)
        ws.merge_cells(start_row=sig_row + 1, start_column=1, end_row=sig_row + 1, end_column=mid_col)
        
        c1_line = ws.cell(row=sig_row, column=1, value="_________________________________________")
        c1_line.font = font_td_bold
        c1_line.alignment = Alignment(horizontal="center", vertical="center")
        
        c1_text = ws.cell(row=sig_row + 1, column=1, value=data.signatures[0])
        c1_text.font = font_footer
        c1_text.alignment = Alignment(horizontal="center", vertical="center")

        # Signature 2 (Right Half)
        for c in range(mid_col + 1, num_cols + 1):
            ws.cell(row=sig_row, column=c)
            ws.cell(row=sig_row + 1, column=c)
        ws.merge_cells(start_row=sig_row, start_column=mid_col + 1, end_row=sig_row, end_column=num_cols)
        ws.merge_cells(start_row=sig_row + 1, start_column=mid_col + 1, end_row=sig_row + 1, end_column=num_cols)
        
        c2_line = ws.cell(row=sig_row, column=mid_col + 1, value="_________________________________________")
        c2_line.font = font_td_bold
        c2_line.alignment = Alignment(horizontal="center", vertical="center")
        
        c2_text = ws.cell(row=sig_row + 1, column=mid_col + 1, value=data.signatures[1])
        c2_text.font = font_footer
        c2_text.alignment = Alignment(horizontal="center", vertical="center")
    elif data.signatures:
        ws.merge_cells(start_row=sig_row, start_column=1, end_row=sig_row, end_column=num_cols)
        ws.merge_cells(start_row=sig_row + 1, start_column=1, end_row=sig_row + 1, end_column=num_cols)
        c_line = ws.cell(row=sig_row, column=1, value="_________________________________________")
        c_line.font = font_td_bold
        c_line.alignment = Alignment(horizontal="center", vertical="center")
        c_text = ws.cell(row=sig_row + 1, column=1, value=data.signatures[0])
        c_text.font = font_footer
        c_text.alignment = Alignment(horizontal="center", vertical="center")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
